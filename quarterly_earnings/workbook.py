from __future__ import annotations
import argparse
import csv
import json
import logging
import os
import re
import shutil
import time
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from xml.etree import ElementTree
from zoneinfo import ZoneInfo
import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill

from .constants import TEMPLATE_PATH, QUARTER_LABELS as L
from .models import CalcResult, Disclosure
from .util import dq, is_loss_turnaround, is_turnaround
from .stats import count_reason

def build_workbook(
    output_path: Path,
    args: argparse.Namespace,
    run_started: datetime,
    raw_rows: list[dict[str, Any]],
    calc_results: list[CalcResult],
    passed_results: list[CalcResult],
    excluded_rows: list[dict[str, Any]],
    quality_rows: list[dict[str, Any]],
    quarter_rows: list[dict[str, Any]],
) -> None:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE_PATH}")
    shutil.copy2(TEMPLATE_PATH, output_path)
    wb = openpyxl.load_workbook(output_path)
    ws = wb.worksheets[0]
    ws.title = "Main"
    clear_main_values(ws)
    write_main_rows(ws, passed_results)

    for sheet_name in ["Summary", "Raw_Disclosures", "Calculation_Detail", "Quarter_Data", "Excluded", "DataQuality"]:
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]

    write_sheet(wb, "Summary", summary_rows(args, run_started, raw_rows, passed_results, excluded_rows, quality_rows))
    write_sheet(wb, "Raw_Disclosures", raw_rows)
    write_sheet(wb, "Calculation_Detail", [calculation_row(result) for result in calc_results])
    write_sheet(wb, "Quarter_Data", quarter_rows)
    write_sheet(wb, "Excluded", excluded_rows)
    write_sheet(wb, "DataQuality", quality_rows or [dq("", "DataQuality", "OK", "확인 필요 항목 없음", False, "")])

    wb.save(output_path)

def clear_main_values(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
    max_row = max(ws.max_row, 3)
    for row in range(3, max_row + 1):
        for col in range(1, 14):
            ws.cell(row=row, column=col).value = None

def write_main_rows(ws: openpyxl.worksheet.worksheet.Worksheet, results: list[CalcResult]) -> None:
    for idx, result in enumerate(results, start=3):
        if idx > ws.max_row:
            copy_row_style(ws, 3, idx)
        values = [
            result.disclosure.corp_name,
            result.business_model,
            result.market_cap_eok if result.market_cap_eok is not None else "확인 필요",
            result.h1_revenue if result.h1_revenue is not None else "확인 필요",
            result.q2_operating_income if result.q2_operating_income is not None else "확인 필요",
            result.h1_operating_income if result.h1_operating_income is not None else "확인 필요",
            round(result.per, 1) if result.per is not None else "N/A",
            display_metric(result.qoq_revenue_growth),
            display_op_change(result.q2_operating_income, result.q1_operating_income, result.qoq_op_growth),
            display_metric(result.yoy_revenue_growth),
            display_op_change(result.q2_operating_income, result.y2025_q2_operating_income, result.yoy_op_growth),
            display_metric(result.yoy_cum_revenue_growth),
            # 누적 영익도 부호가 바뀌면 흑전/적전으로 표시 (QoQ/YoY 컬럼과 동일 규칙)
            display_op_change(result.h1_operating_income, result.y2025_h1_operating_income, result.yoy_cum_op_growth),
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=idx, column=col).value = value

def copy_row_style(ws: openpyxl.worksheet.worksheet.Worksheet, source_row: int, target_row: int) -> None:
    for col in range(1, 14):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = source._style.copy()
        target.number_format = source.number_format
        target.alignment = source.alignment.copy()
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height

def write_sheet(wb: openpyxl.Workbook, name: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(name)
    if not rows:
        rows = [{"상태": "데이터 없음"}]
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    style_header(ws)
    autofit(ws)

def style_header(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
    fill = PatternFill("solid", fgColor="D9EAD3")
    font = Font(name="맑은 고딕", size=10, bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

def autofit(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
    for column in ws.columns:
        max_len = 0
        column_letter = column[0].column_letter
        for cell in column:
            max_len = max(max_len, len(str(cell.value or "")))
        ws.column_dimensions[column_letter].width = min(max(max_len + 2, 10), 45)

def summary_rows(
    args: argparse.Namespace,
    run_started: datetime,
    raw_rows: list[dict[str, Any]],
    passed_results: list[CalcResult],
    excluded_rows: list[dict[str, Any]],
    quality_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    type_counts: dict[str, int] = {}
    for result in passed_results:
        type_counts[result.disclosure.report_type] = type_counts.get(result.disclosure.report_type, 0) + 1
    return [
        {"항목": "실행일", "값": run_started.strftime("%Y-%m-%d %H:%M:%S")},
        {"항목": "검색 기간", "값": f"{args.start} ~ {args.end}"},
        {"항목": "대상 시장", "값": args.market},
        {"항목": "전체 공시 수", "값": len(raw_rows)},
        {"항목": "조건 부합 종목 수", "값": len(passed_results)},
        {"항목": "순수 보통주 제외 처리 종목 수", "값": len(excluded_rows)},
        {"항목": "공시 유형별 개수", "값": json.dumps(type_counts, ensure_ascii=False)},
        {"항목": "QoQ 매출 15% 이상 종목 수", "값": count_reason(passed_results, "QoQ 매출")},
        {"항목": "QoQ 영익 15% 이상 종목 수", "값": count_reason(passed_results, "QoQ 영익")},
        {"항목": "YoY 매출 15% 이상 종목 수", "값": count_reason(passed_results, "YoY 매출")},
        {"항목": "YoY 영익 15% 이상 종목 수", "값": count_reason(passed_results, "YoY 영익")},
        {"항목": "흑자전환 종목 수", "값": sum(1 for r in passed_results if r.qoq_turnaround or r.yoy_turnaround)},
        {"항목": "PER 계산 기준", "값": "시가총액 / 누적 영업이익"},
        {"항목": "확인 필요 항목 수", "값": sum(1 for row in quality_rows if row.get("확인 필요 여부") == "Y")},
    ]

def calculation_row(result: CalcResult) -> dict[str, Any]:
    return {
        "종목명": result.disclosure.corp_name,
        "종목코드": result.disclosure.stock_code,
        "공시 유형": result.disclosure.report_type,
        f"{L['cur']} 매출": display_value(result.q2_revenue),
        f"{L['cur']} 영업이익": display_value(result.q2_operating_income),
        f"{L['cum']} 매출": display_value(result.h1_revenue),
        f"{L['cum']} 영업이익": display_value(result.h1_operating_income),
        f"{L['prev']} 매출": display_value(result.q1_revenue),
        f"{L['prev']} 영업이익": display_value(result.q1_operating_income),
        f"{L['yoy']} 매출": display_value(result.y2025_q2_revenue),
        f"{L['yoy']} 영업이익": display_value(result.y2025_q2_operating_income),
        f"{L['yoycum']} 매출": display_value(result.y2025_h1_revenue),
        f"{L['yoycum']} 영업이익": display_value(result.y2025_h1_operating_income),
        "QoQ 매출 성장률": display_metric(result.qoq_revenue_growth),
        "QoQ 영익 성장률": display_op_change(result.q2_operating_income, result.q1_operating_income, result.qoq_op_growth),
        "YoY 매출 성장률": display_metric(result.yoy_revenue_growth),
        "YoY 영익 성장률": display_op_change(result.q2_operating_income, result.y2025_q2_operating_income, result.yoy_op_growth),
        "YoY누적 매출 성장률": display_metric(result.yoy_cum_revenue_growth),
        "YoY누적 영익 성장률": display_op_change(result.h1_operating_income, result.y2025_h1_operating_income, result.yoy_cum_op_growth),
        "흑자/적자전환 여부": turnaround_label(result),
        "조건 통과 사유": ", ".join(result.pass_reasons),
        "시가총액": display_value(result.market_cap_eok),
        "PER 계산 기준 영업이익": display_value(result.per_basis_op_income),
        "PER": round(result.per, 1) if result.per is not None else "N/A",
        "계산 방식": result.calculation_method,
        "출처 링크": result.disclosure.link,
    }

def disclosure_to_row(item: Disclosure) -> dict[str, Any]:
    return {
        "접수번호": item.rcp_no,
        "종목명": item.corp_name,
        "종목코드": item.stock_code,
        "공시일": item.rcept_dt,
        "공시명": item.report_nm,
        "공시 유형": item.report_type,
        "시장": item.corp_cls,
        "공시 링크": item.link,
    }

def make_summary(
    args: argparse.Namespace,
    run_started: datetime,
    raw_rows: list[dict[str, Any]],
    passed_results: list[CalcResult],
    excluded_rows: list[dict[str, Any]],
    quality_rows: list[dict[str, Any]],
) -> dict[str, Any]:
    return {
        "run_started": run_started.isoformat(),
        "start": args.start,
        "end": args.end,
        "quarter": args.quarter,
        "market": args.market,
        "raw_disclosures": len(raw_rows),
        "passed_stocks": len(passed_results),
        "excluded_stocks": len(excluded_rows),
        "turnaround_stocks": sum(1 for r in passed_results if r.qoq_turnaround or r.yoy_turnaround),
        "data_quality_items": len(quality_rows),
        "per_basis": "market_cap_eok / cumulative_operating_income_eok",
    }

def print_run_summary(
    args: argparse.Namespace,
    output_path: Path,
    raw_rows: list[dict[str, Any]],
    passed_results: list[CalcResult],
    excluded_rows: list[dict[str, Any]],
    quality_rows: list[dict[str, Any]],
) -> None:
    type_counts: dict[str, int] = {}
    for result in passed_results:
        type_counts[result.disclosure.report_type] = type_counts.get(result.disclosure.report_type, 0) + 1
    print(json.dumps(
        {
            "output": str(output_path),
            "search_period": f"{args.start} ~ {args.end}",
            "raw_disclosures": len(raw_rows),
            "passed_stocks": len(passed_results),
            "excluded_stocks": len(excluded_rows),
            "passed_by_report_type": type_counts,
            "turnaround_stocks": sum(1 for r in passed_results if r.qoq_turnaround or r.yoy_turnaround),
            "per_basis": "시가총액 / 누적 영업이익",
            "data_quality_items": len(quality_rows),
        },
        ensure_ascii=False,
        indent=2,
    ))

def display_metric(value: float | None) -> float | str:
    return round(value, 2) if value is not None else "확인 필요"

def display_op_change(current: float | None, base: float | None, growth_value: float | None) -> float | str:
    if is_turnaround(current, base):
        return "흑전"
    if is_loss_turnaround(current, base):
        return "적전"
    return display_metric(growth_value)

def display_value(value: float | None) -> float | str:
    return round(value, 2) if value is not None else "확인 필요"

def turnaround_label(result: CalcResult) -> str:
    labels = []
    if result.qoq_turnaround:
        labels.append("전분기 흑전")
    if result.yoy_turnaround:
        labels.append("전년동기 흑전")
    if result.qoq_loss_turnaround:
        labels.append("전분기 적전")
    if result.yoy_loss_turnaround:
        labels.append("전년동기 적전")
    return ", ".join(labels) if labels else ""
